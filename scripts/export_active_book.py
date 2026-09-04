"""Export True Active policies from Master Book v3 (Ignoring CommBook for Cancellations)."""
import sys, os, csv, logging
from datetime import datetime, date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "worker"))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("export_active")

TODAY = date.today()

def parse_date(s):
    if not s: return None
    try: return datetime.strptime(s.strip()[:10], "%m/%d/%Y").date()
    except:
        try: return datetime.strptime(s.strip()[:10], "%Y-%m-%d").date()
        except: return None

def main():
    base_dir = os.path.join(os.path.dirname(__file__), "..")
    exports_dir = os.path.join(base_dir, "exports")

    in_csv = os.path.join(exports_dir, "CFP_Master_Book_v3_2026-09-03.csv")
    date_stamp = TODAY.strftime("%Y-%m-%d")
    out_xlsx = os.path.join(exports_dir, f"CFP_Active_Book_{date_stamp}.xlsx")
    out_csv = os.path.join(exports_dir, f"CFP_Active_Book_{date_stamp}.csv")

    with open(in_csv, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames
        rows = list(reader)

    active_rows = []
    
    for r in rows:
        status = r.get("Policy Status", "").upper()
        priority = r.get("Priority Status", "").upper()
        # We NO LONGER check `reason` for PCAN/FCAN
        cancel = r.get("Cancellation Reason", "").upper()
        eff = parse_date(r.get("Effective Date"))
        exp = parse_date(r.get("Expiration Date"))
        
        # 1. Exclude explicit cancellations (IGNORING COMMBOOK REASON)
        if "CANCELLED" in status or "CANCELLED" in priority or cancel:
            continue
            
        is_active = False
        
        # 2. Check dates and tracker priority
        if exp:
            if exp >= TODAY:
                is_active = True
            elif "DONE RENEWAL" in priority:
                is_active = True
        elif eff:
            if (TODAY - eff).days <= 365:
                is_active = True
                
        if is_active:
            r["Policy Status"] = "Active"
            active_rows.append(r)

    logger.info("Found %d True Active policies (ignoring CommBook)", len(active_rows))

    # --- CSV ---
    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(active_rows)
        
    # --- Excel ---
    wb = Workbook()
    ws = wb.active
    ws.title = "Active Policies"
    
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        
    data_font = Font(name="Calibri", size=10)
    alt_fill = PatternFill(start_color="F2F6FC", end_color="F2F6FC", fill_type="solid")
    
    for row_idx, row_dict in enumerate(active_rows, 2):
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_dict.get(header, ""))
            cell.font = data_font
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill
                
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(active_rows) + 1}"
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 15
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['J'].width = 35
    ws.column_dimensions['S'].width = 40
    
    # Add Total row
    prem_col_idx = headers.index("Annual Premium") + 1
    total_row = len(active_rows) + 3
    
    total_prem = 0
    for r in active_rows:
        p_str = r.get("Annual Premium", "").replace("$", "").replace(",", "")
        if p_str:
            try: total_prem += float(p_str)
            except: pass

    label_cell = ws.cell(row=total_row, column=prem_col_idx - 1, value="TOTAL PREMIUM:")
    label_cell.font = Font(name="Calibri", bold=True, size=11)
    label_cell.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    label_cell.border = thin_border

    val_cell = ws.cell(row=total_row, column=prem_col_idx, value=f"${total_prem:,.2f}")
    val_cell.font = Font(name="Calibri", bold=True, size=11)
    val_cell.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    val_cell.border = thin_border

    wb.save(out_xlsx)
    logger.info("Saved True Active export with total premium %s", total_prem)

if __name__ == "__main__":
    main()
