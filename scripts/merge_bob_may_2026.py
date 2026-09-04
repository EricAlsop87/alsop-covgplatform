"""
Merge #3: Integrate Alsop BOB May 2026 into Master Book.
==========================================================
- Fuzzy matches on Property Address OR (3-char Last Name + Dates).
- Populates empty coverage columns in Master Book.
- Keeps newer Effective/Expiration dates.
- Outputs unmatched BOB rows to a separate tab in Excel.
"""

import sys, os, re, csv, logging
from datetime import datetime, date
from collections import Counter

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "worker"))
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("merge3")

TODAY = date.today()

def normalize_pn(pn: str) -> str:
    return re.sub(r'\s+', ' ', pn.strip().upper())

def norm_addr(addr):
    if not addr: return ""
    s = str(addr).upper()
    s = re.sub(r'[^A-Z0-9 ]', '', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s

def safe_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%m/%d/%Y")
    if isinstance(val, date):
        return val.strftime("%m/%d/%Y")
    if isinstance(val, bool):
        return "Yes" if val else "No"
    if isinstance(val, float):
        if val == int(val):
            return str(int(val))
        return str(val)
    s = str(val).strip()
    if s in ("None", "-"):
        return ""
    return s

def parse_date(val):
    if isinstance(val, (datetime, date)):
        return val
    s = str(val).strip()
    if len(s) >= 10:
        try:
            return datetime.strptime(s[:10], "%Y-%m-%d")
        except:
            try:
                return datetime.strptime(s[:10], "%m/%d/%Y")
            except:
                pass
    return None

def coalesce(*values) -> str:
    for v in values:
        s = v if isinstance(v, str) else safe_str(v)
        if s and s.strip():
            return s.strip()
    return ""

def load_master_csv(path: str) -> tuple[list[str], list[dict]]:
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames
        rows = list(reader)
    logger.info("Loaded %d master rows", len(rows))
    return headers, rows

def load_bob(xlsx_path: str):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws_d = wb['Dwelling']
    raw_d = list(ws_d.iter_rows(values_only=True))
    headers_d = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(raw_d[0])]
    data_d = [r for r in raw_d[1:] if any(c is not None for c in r)]
    
    ws_c = wb['Commercial']
    raw_c = list(ws_c.iter_rows(values_only=True))
    headers_c = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(raw_c[0])]
    data_c = [r for r in raw_c[1:] if any(c is not None for c in r)]
    
    wb.close()
    
    # Convert to dicts
    dwellings = []
    for row in data_d:
        d = {h: row[i] if i < len(row) else None for i, h in enumerate(headers_d)}
        d['__sheet'] = 'Dwelling'
        dwellings.append(d)
        
    commercials = []
    for row in data_c:
        d = {h: row[i] if i < len(row) else None for i, h in enumerate(headers_c)}
        d['__sheet'] = 'Commercial'
        commercials.append(d)
        
    logger.info("Loaded %d Dwelling and %d Commercial BOB rows", len(dwellings), len(commercials))
    return dwellings + commercials

def main():
    base_dir = os.path.join(os.path.dirname(__file__), "..")
    exports_dir = os.path.join(base_dir, "exports")

    master_csv = os.path.join(exports_dir, "CFP_Master_Book_v2_2026-09-03.csv")
    bob_xlsx = os.path.join(exports_dir, "Alsop BOB May 2026.xlsx")
    date_stamp = TODAY.strftime("%Y-%m-%d")
    out_xlsx = os.path.join(exports_dir, f"CFP_Master_Book_v3_{date_stamp}.xlsx")
    out_csv = os.path.join(exports_dir, f"CFP_Master_Book_v3_{date_stamp}.csv")

    master_headers, master_rows = load_master_csv(master_csv)
    bob_rows = load_bob(bob_xlsx)
    
    # Build match indices for Master
    master_by_addr = {}
    master_by_3char = {}
    
    for i, r in enumerate(master_rows):
        addr = norm_addr(r.get('Property Address', ''))
        if addr:
            master_by_addr.setdefault(addr, []).append(i)
            
        lname = str(r.get('Last Name', '')).strip().upper()[:3]
        eff = str(r.get('Effective Date', '')).strip()
        exp = str(r.get('Expiration Date', '')).strip()
        if lname and eff and exp:
            key = f"{lname}|{eff}|{exp}"
            master_by_3char.setdefault(key, []).append(i)
            
    # Attempt to match BOB rows to Master rows
    bob_matched_to_master = {} # bob_index -> master_index
    master_matched_from_bob = {} # master_index -> list of bob_rows
    unmatched_bob = []
    
    for bob_idx, b in enumerate(bob_rows):
        is_comm = b['__sheet'] == 'Commercial'
        
        # Address matching
        if is_comm:
            addr_val = safe_str(b.get('Risk Address', ''))
            city_val = safe_str(b.get('Risk City', ''))
        else:
            addr_val = safe_str(b.get('Risk Address', ''))
            city_val = safe_str(b.get('Risk City', ''))
            
        addr = norm_addr(addr_val)
        full_addr = norm_addr(f"{addr_val} {city_val}")
        
        # Date + Name matching
        lname = safe_str(b.get('Last Name', '')).upper()[:3]
        eff_date = parse_date(b.get('Effective Date', ''))
        exp_date = parse_date(b.get('Expire Date', ''))
        
        eff_str = eff_date.strftime("%m/%d/%Y") if eff_date else ""
        exp_str = exp_date.strftime("%m/%d/%Y") if exp_date else ""
        date_key = f"{lname}|{eff_str}|{exp_str}"
        
        matched_idx = None
        
        # 1. Exact address match
        if addr and addr in master_by_addr and len(master_by_addr[addr]) == 1:
            matched_idx = master_by_addr[addr][0]
            
        # 2. Date + Name match
        if matched_idx is None and eff_str and exp_str and date_key in master_by_3char:
            if len(master_by_3char[date_key]) == 1:
                matched_idx = master_by_3char[date_key][0]
                
        # 3. Fuzzy address match
        if matched_idx is None and addr:
            for ma, indices in master_by_addr.items():
                if (addr in ma or ma in addr) and len(indices) == 1:
                    matched_idx = indices[0]
                    break
                    
        if matched_idx is not None:
            bob_matched_to_master[bob_idx] = matched_idx
            master_matched_from_bob.setdefault(matched_idx, []).append(b)
        else:
            unmatched_bob.append(b)
            
    logger.info("Matched %d BOB rows to Master", len(bob_matched_to_master))
    logger.info("Unmatched BOB rows: %d", len(unmatched_bob))
    
    # Process Master rows and enrich
    enriched_master = []
    enriched_count = 0
    
    def format_money(val):
        s = safe_str(val)
        if not s or s == "0": return ""
        try:
            num = float(s.replace(",", "").replace("$", ""))
            return f"${int(num):,}" if num == int(num) else f"${num:,.2f}"
        except:
            return s
            
    for i, m_row in enumerate(master_rows):
        b_list = master_matched_from_bob.get(i, [])
        if not b_list:
            enriched_master.append(m_row)
            continue
            
        enriched_count += 1
        # If multiple BOB rows matched one Master row, just take the first one
        b = b_list[0]
        is_comm = b['__sheet'] == 'Commercial'
        
        new_row = m_row.copy()
        
        # Append data source
        ds = new_row.get("Data Source", "")
        if "Alsop BOB" not in ds:
            new_row["Data Source"] = ds + " + Alsop BOB" if ds else "Alsop BOB"
            
        # Compare dates and keep newer
        m_eff = parse_date(new_row.get("Effective Date"))
        b_eff = parse_date(b.get('Effective Date'))
        if b_eff and (not m_eff or b_eff > m_eff):
            new_row["Effective Date"] = b_eff.strftime("%m/%d/%Y")
            
        m_exp = parse_date(new_row.get("Expiration Date"))
        b_exp = parse_date(b.get('Expire Date'))
        if b_exp and (not m_exp or b_exp > m_exp):
            new_row["Expiration Date"] = b_exp.strftime("%m/%d/%Y")
            # Update Policy Status based on new date
            if new_row.get("Policy Status") not in ("Active", "Cancelled"): # Don't override explicit human status
                new_row["Policy Status"] = "Active" if b_exp >= datetime.now() else "Expired"
                
        # Enrich coverage limits (only overwrite if empty)
        if not new_row.get("Coverage A (Dwelling)"):
            if is_comm: new_row["Coverage A (Dwelling)"] = format_money(b.get('Dwelling coverage limit'))
            else: new_row["Coverage A (Dwelling)"] = format_money(b.get('Dwelling Coverage Limit'))
            
        if not new_row.get("Coverage B (Other Structures)"):
            if is_comm: new_row["Coverage B (Other Structures)"] = format_money(b.get('Other structures total limit'))
            else: new_row["Coverage B (Other Structures)"] = format_money(b.get('Other Structures Total Limit'))
            
        if not new_row.get("Coverage C (Personal Property)"):
            if is_comm: new_row["Coverage C (Personal Property)"] = format_money(b.get('BPP'))
            else: new_row["Coverage C (Personal Property)"] = format_money(b.get('Personal Property Amount'))
            
        if not new_row.get("Coverage D (Fair Rental Value)") and not is_comm:
            new_row["Coverage D (Fair Rental Value)"] = format_money(b.get('Fair Rental Value'))
            
        if not new_row.get("Ordinance or Law") and not is_comm:
            new_row["Ordinance or Law"] = format_money(b.get('Ordinance of Law Coverage'))
            
        if not new_row.get("Debris Removal") and not is_comm:
            new_row["Debris Removal"] = format_money(b.get('Debris Removal'))
            
        if not new_row.get("Deductible") and not is_comm:
            new_row["Deductible"] = format_money(b.get('Deductible'))
            
        if not new_row.get("Construction Type"):
            new_row["Construction Type"] = safe_str(b.get('Construction'))
            
        if not new_row.get("Year Built") and not is_comm:
            new_row["Year Built"] = safe_str(b.get('Construction Year'))
            
        if not new_row.get("Occupancy") and not is_comm:
            new_row["Occupancy"] = safe_str(b.get('Occupancy'))
            
        enriched_master.append(new_row)
        
    logger.info("Enriched %d Master rows", enriched_count)
    
    # Write output CSV
    logger.info("Writing CSV to %s ...", out_csv)
    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=master_headers)
        writer.writeheader()
        writer.writerows(enriched_master)
        
    # Write output Excel with 2 tabs
    logger.info("Writing Excel to %s ...", out_xlsx)
    wb = Workbook()
    
    # Tab 1: Master Book
    ws1 = wb.active
    ws1.title = "Master Book"
    
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    
    for col_idx, header in enumerate(master_headers, 1):
        cell = ws1.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        
    data_font = Font(name="Calibri", size=10)
    alt_fill = PatternFill(start_color="F2F6FC", end_color="F2F6FC", fill_type="solid")
    
    for row_idx, row_dict in enumerate(enriched_master, 2):
        for col_idx, header in enumerate(master_headers, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=row_dict.get(header, ""))
            cell.font = data_font
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill
                
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(master_headers))}{len(enriched_master) + 1}"
    for col_idx in range(1, len(master_headers) + 1):
        ws1.column_dimensions[get_column_letter(col_idx)].width = 15
    ws1.column_dimensions['A'].width = 22 # Policy Number
    ws1.column_dimensions['C'].width = 20 # Carrier Name
    ws1.column_dimensions['J'].width = 35 # Full Named Insured
    ws1.column_dimensions['S'].width = 40 # Property Address
    
    # Tab 2: Unmatched BOB
    ws2 = wb.create_sheet(title="Unmatched BOB Policies")
    
    if unmatched_bob:
        # Determine all unique headers across dwelling and commercial
        unmatched_headers = []
        for b in unmatched_bob:
            for k in b.keys():
                if k not in unmatched_headers:
                    unmatched_headers.append(k)
                    
        # Put __sheet first
        if '__sheet' in unmatched_headers:
            unmatched_headers.remove('__sheet')
            unmatched_headers.insert(0, '__sheet')
            
        for col_idx, header in enumerate(unmatched_headers, 1):
            cell = ws2.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
            
        for row_idx, row_dict in enumerate(unmatched_bob, 2):
            for col_idx, header in enumerate(unmatched_headers, 1):
                val = row_dict.get(header)
                if isinstance(val, (datetime, date)):
                    val = val.strftime("%m/%d/%Y")
                cell = ws2.cell(row=row_idx, column=col_idx, value=safe_str(val))
                cell.font = data_font
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = alt_fill
                    
        ws2.freeze_panes = "A2"
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(unmatched_headers))}{len(unmatched_bob) + 1}"
        for col_idx in range(1, len(unmatched_headers) + 1):
            ws2.column_dimensions[get_column_letter(col_idx)].width = 18
            
    wb.save(out_xlsx)
    logger.info("  Excel saved")
    
    print("\n" + "=" * 70)
    print("  MASTER BOOK v3 — VERIFICATION SUMMARY")
    print("=" * 70)
    print(f"  Total Master Rows:     {len(enriched_master)}")
    print(f"  Master Rows Enriched:  {enriched_count}")
    print(f"  Unmatched BOB Rows:    {len(unmatched_bob)} (saved to separate tab)")
    
    # Coverage completeness check
    cov_a_count = sum(1 for r in enriched_master if r.get('Coverage A (Dwelling)'))
    deduct_count = sum(1 for r in enriched_master if r.get('Deductible'))
    
    print("\n  DATA ENRICHMENT RESULTS:")
    print(f"    Coverage A (Dwelling) filled: {cov_a_count} ({cov_a_count/len(enriched_master)*100:.1f}%)")
    print(f"    Deductible filled:            {deduct_count} ({deduct_count/len(enriched_master)*100:.1f}%)")
    print("=" * 70)

if __name__ == "__main__":
    main()
